from __future__ import annotations

import os
import subprocess
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

import formualizer as fz


@pytest.mark.parametrize("size", [2, 5])
@pytest.mark.parametrize("backend", ["python", "cli"])
def test_openpyxl_resave_becomes_fixed_extent(tmp_path, size, backend):
    openpyxl = pytest.importorskip("openpyxl")
    cli = os.environ.get("FORMUALIZER_CLI")
    if backend == "cli" and not cli:
        pytest.skip("set FORMUALIZER_CLI to test the candidate executable")
    path = tmp_path / "roundtrip.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 3
    workbook.active["B1"] = "=SEQUENCE(A1)"
    workbook.active["D1"] = "=SUM(B1:B3)"
    workbook.active["E1"] = "=SUM(B1#)"
    workbook.save(path)

    def recalc():
        if backend == "python":
            fz.recalculate_xlsx_file(str(path))
        else:
            result = subprocess.run([cli, "recalc", str(path)], capture_output=True)
            assert result.returncode == 0, result.stderr.decode()

    recalc()
    workbook = openpyxl.load_workbook(path)
    workbook.active["A1"] = size
    workbook.save(path)
    recalc()
    cached = openpyxl.load_workbook(path, data_only=True).active
    assert [cached[f"B{row}"].value for row in range(1, 4)] == (
        [1, 2, "#N/A"] if size == 2 else [1, 2, 3]
    )
    assert cached["D1"].value == ("#N/A" if size == 2 else 6)
    assert cached["E1"].value == "#REF!"
    formulas = openpyxl.load_workbook(path).active
    assert formulas["B1"].value.ref == "B1:B3"
    with ZipFile(path) as archive:
        assert "xl/metadata.xml" not in archive.namelist()
        xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cell = xml.find('.//{*}c[@r="B1"]')
        assert "cm" not in cell.attrib


@pytest.mark.parametrize("backend", ["python", "cli"])
def test_array_if_cse_and_dynamic_roundtrip(tmp_path, backend):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.worksheet.formula import ArrayFormula

    cli = os.environ.get("FORMUALIZER_CLI")
    if backend == "cli" and not cli:
        pytest.skip("set FORMUALIZER_CLI to test the candidate executable")
    path = tmp_path / "array-if.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in range(1, 4):
        sheet.cell(row, 1, row)
    sheet["B1"] = ArrayFormula(ref="B1", text="=SUM(IF(A1:A3>0,A1:A3))")
    sheet["C1"] = "=IF(A1:A3>1,A1:A3,0)"
    workbook.save(path)
    for _ in range(2):
        if backend == "python":
            fz.recalculate_xlsx_file(str(path))
        else:
            result = subprocess.run([cli, "recalc", str(path)], capture_output=True)
            assert result.returncode == 0, result.stderr.decode()
        cached = openpyxl.load_workbook(path, data_only=True)
        assert cached.active["B1"].value == 6
        assert [cached.active[f"C{row}"].value for row in range(1, 4)] == [0, 2, 3]
        cached.close()
    formulas = openpyxl.load_workbook(path)
    assert formulas.active["B1"].value.ref == "B1"
    assert formulas.active["C1"].value.ref == "C1:C3"
    formulas.close()


def fixture_xlsx(*, formula: bool = True) -> bytes:
    worksheet = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>2</v></c>'
        + (
            '<c r="C1" t="str"><f>A1+B1</f><v>stale</v></c>'
            if formula
            else '<c r="C1"><v>3</v></c>'
        )
        + "</row></sheetData></worksheet>"
    )
    members = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": worksheet,
    }
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        for name, body in members.items():
            archive.writestr(name, body)
    return out.getvalue()


def calculation_name_fixture(definitions: str, formula: str) -> bytes:
    out = BytesIO()
    with (
        ZipFile(BytesIO(fixture_xlsx())) as source,
        ZipFile(out, "w", ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            body = source.read(name).decode()
            if name == "xl/workbook.xml":
                body = body.replace(
                    "</workbook>",
                    f"<definedNames>{definitions}</definedNames></workbook>",
                )
            elif name == "xl/worksheets/sheet1.xml":
                body = body.replace("A1+B1", formula)
            target.writestr(name, body)
    return out.getvalue()


@pytest.mark.parametrize(
    "definition,formula,expected,kind",
    [
        ("0.07", "Rate*2", "0.14", None),
        ("Sheet1!$A$1*2", "Rate", "2", None),
        ("TRUE", "Rate", "1", "b"),
        ("&quot;original text&quot;", "Rate", "original text", "str"),
        ("#N/A", "Rate", "#N/A", "e"),
    ],
)
def test_calculation_names_preserve_metadata_and_typed_results(
    tmp_path, definition, formula, expected, kind
):
    payload = calculation_name_fixture(
        f'<definedName name="Rate">{definition}</definedName>', formula
    )
    result = fz.recalculate_xlsx_bytes(payload)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = ET.fromstring(worksheet_xml(result["bytes"])).find('.//s:c[@r="C1"]', ns)
    assert cell.get("t") == kind
    assert cell.findtext("s:v", namespaces=ns) == expected
    with (
        ZipFile(BytesIO(payload)) as before,
        ZipFile(BytesIO(result["bytes"])) as after,
    ):
        assert before.read("xl/workbook.xml") == after.read("xl/workbook.xml")
    assert fz.recalculate_xlsx_bytes(result["bytes"])["bytes"] == result["bytes"]
    source, destination = tmp_path / "source.xlsx", tmp_path / "out.xlsx"
    source.write_bytes(payload)
    fz.recalculate_xlsx_file(str(source), output=str(destination))
    assert source.read_bytes() == payload
    assert destination.read_bytes() == result["bytes"]


@pytest.mark.parametrize(
    "definitions",
    [
        '<definedName name="Rate">Sheet1!$A$1,Sheet1!$B$1</definedName>',
        '<definedName name="Rate">Sheet1!A1*2</definedName>',
        '<definedName name="Rate">Rate</definedName>',
    ],
)
def test_unsupported_calculation_names_do_not_publish(tmp_path, definitions):
    payload = calculation_name_fixture(definitions, "Rate*2")
    with pytest.raises(OSError, match="recalculate XLSX failed"):
        fz.recalculate_xlsx_bytes(payload)
    source, destination = tmp_path / "source.xlsx", tmp_path / "out.xlsx"
    source.write_bytes(payload)
    destination.write_bytes(b"keep destination")
    with pytest.raises(OSError, match="recalculate XLSX failed"):
        fz.recalculate_xlsx_file(str(source), output=str(destination))
    assert source.read_bytes() == payload
    assert destination.read_bytes() == b"keep destination"


def worksheet_xml(payload: bytes) -> str:
    with ZipFile(BytesIO(payload)) as archive:
        return archive.read("xl/worksheets/sheet1.xml").decode()


def test_recalculate_xlsx_bytes_updates_typed_cache_and_returns_bytes():
    result = fz.recalculate_xlsx_bytes(fixture_xlsx())

    assert isinstance(result["bytes"], bytes)
    assert result["summary"]["status"] == "success"
    assert result["formula_cells"] == result["cache_cells_changed"] == 1
    assert result["worksheet_parts_changed"] == 1
    xml = worksheet_xml(result["bytes"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = ET.fromstring(xml).find('.//s:c[@r="C1"]', ns)
    assert cell is not None and cell.get("t") is None
    assert cell.findtext("s:f", namespaces=ns) == "A1+B1"
    assert cell.findtext("s:v", namespaces=ns) == "3"
    repeated = fz.recalculate_xlsx_bytes(result["bytes"])
    assert repeated["bytes"] == result["bytes"]
    assert repeated["cache_cells_changed"] == 0


def test_recalculate_xlsx_bytes_no_formula_is_a_noop():
    payload = fixture_xlsx(formula=False)
    result = fz.recalculate_xlsx_bytes(payload)

    assert result["bytes"] == payload
    assert result["formula_cells"] == 0
    assert result["cache_cells_changed"] == 0
    assert result["worksheet_parts_changed"] == 0


def test_recalculate_xlsx_bytes_surfaces_core_rejection():
    with pytest.raises(OSError, match="recalculate XLSX failed"):
        fz.recalculate_xlsx_bytes(b"not an XLSX package")


def test_recalculate_xlsx_file_failure_leaves_existing_destination_unchanged(
    tmp_path: Path,
):
    source = tmp_path / "invalid.xlsx"
    destination = tmp_path / "destination.xlsx"
    source.write_bytes(b"not an XLSX package")
    destination.write_bytes(b"must remain unchanged")

    with pytest.raises(OSError, match="recalculate XLSX failed"):
        fz.recalculate_xlsx_file(str(source), output=str(destination))

    assert destination.read_bytes() == b"must remain unchanged"


def test_error_reasons_and_unknown_functions_are_reported():
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "=SPDVOL(1)"
    sheet["A2"] = "=SPDVOL(2)"
    sheet["A3"] = "=_xll.EURO(1)"
    sheet["A4"] = "=NoSuchName"
    sheet["A5"] = "=A1+1"
    buffer = BytesIO()
    workbook.save(buffer)
    result = fz.recalculate_xlsx_bytes(buffer.getvalue(), error_location_limit=10)
    summary = result["summary"]
    name = summary["error_summary"]["#NAME?"]
    assert dict(zip(name["locations"], name["messages"], strict=True)) == {
        "Sheet!A1": "Unknown function: SPDVOL",
        "Sheet!A2": "Unknown function: SPDVOL",
        "Sheet!A3": "Unknown function: _xll.EURO",
        "Sheet!A4": "Undefined name: NoSuchName",
        "Sheet!A5": None,
    }
    assert summary["unknown_functions"] == [
        {"name": "SPDVOL", "cells": 2},
        {"name": "_xll.EURO", "cells": 1},
    ]
    truncated = fz.recalculate_xlsx_bytes(buffer.getvalue(), error_location_limit=1)
    assert truncated["summary"]["unknown_functions"] == summary["unknown_functions"]
    assert truncated["summary"]["error_summary"]["#NAME?"]["messages"] == [
        "Unknown function: SPDVOL"
    ]


def _add_calc_chain(path: Path, cells: list[str]) -> None:
    """List `cells` (Sheet1) in a calc chain, as Excel writes for formulas it
    calculated."""
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    entries = "".join(f'<c r="{cell}" i="1"/>' for cell in cells)
    parts["xl/calcChain.xml"] = (
        f'<calcChain xmlns="{main}">{entries}</calcChain>'.encode()
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdChain" Type="http://schemas.openxmlformats.org/'
        b'officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/>'
        b"</Relationships>",
    )
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/calcChain.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_calc_chain_formulas_keep_legacy_intersection(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "legacy.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in range(1, 4):
        sheet[f"A{row}"] = row
        sheet[f"C{row}"] = row
        sheet[f"D{row}"] = f"v{row}"
    sheet["B2"] = "=A1:A3*10"
    sheet["E3"] = "=VLOOKUP(A1:A3,C1:D3,2,FALSE)"
    workbook.save(path)
    agent_written = path.read_bytes()

    # Excel calculated B2 and E3: each intersects its own row.
    _add_calc_chain(path, ["B2", "E3"])
    fz.recalculate_xlsx_file(str(path))
    cached = openpyxl.load_workbook(path, data_only=True).active
    assert cached["B2"].value == 20
    assert cached["E3"].value == "v3"
    assert cached["B3"].value is None
    before = path.read_bytes()
    fz.recalculate_xlsx_file(str(path))
    assert path.read_bytes() == before

    # The same formulas written by openpyxl alone (no calc chain) keep
    # dynamic-array evaluation: B2 spills.
    path.write_bytes(agent_written)
    fz.recalculate_xlsx_file(str(path))
    cached = openpyxl.load_workbook(path, data_only=True).active
    assert [cached[f"B{row}"].value for row in (2, 3, 4)] == [10, 20, 30]


def external_link_fixture(formula: str = "[1]Prices!B2*2") -> bytes:
    """`C1 = [1]Prices!B2*2`, with Excel's cached value 21 for the linked cell."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels = "http://schemas.openxmlformats.org/package/2006/relationships"
    source = BytesIO(fixture_xlsx(formula=False))
    with ZipFile(source) as archive:
        members = {n: archive.read(n).decode() for n in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        '<c r="C1"><v>3</v></c>', f'<c r="C1"><f>{formula}</f><v>0</v></c>'
    )
    members["xl/workbook.xml"] = members["xl/workbook.xml"].replace(
        "</sheets>",
        '</sheets><externalReferences><externalReference r:id="rIdLink1"/></externalReferences>',
    )
    members["xl/_rels/workbook.xml.rels"] = members[
        "xl/_rels/workbook.xml.rels"
    ].replace(
        "</Relationships>",
        f'<Relationship Id="rIdLink1" Type="{office}/externalLink" Target="externalLinks/externalLink1.xml"/></Relationships>',
    )
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        "</Types>",
        '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>',
    )
    members["xl/externalLinks/externalLink1.xml"] = (
        f'<externalLink xmlns="{main}"><externalBook xmlns:r="{office}" r:id="rId1">'
        '<sheetNames><sheetName val="Prices"/></sheetNames><sheetDataSet>'
        '<sheetData sheetId="0"><row r="2"><cell r="B2"><v>21</v></cell></row></sheetData>'
        "</sheetDataSet></externalBook></externalLink>"
    )
    members["xl/externalLinks/_rels/externalLink1.xml.rels"] = (
        f'<Relationships xmlns="{rels}"><Relationship Id="rId1" '
        f'Type="{office}/externalLinkPath" Target="file:///C:/Data/Source.xlsx" TargetMode="External"/></Relationships>'
    )
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        for name, body in members.items():
            archive.writestr(name, body)
    return out.getvalue()


def test_recalculate_xlsx_bytes_reads_cached_external_link_values():
    source = external_link_fixture()
    result = fz.recalculate_xlsx_bytes(source)
    assert result["external_links"] == {
        "links_used": 1,
        "refreshed": False,
        "policy": "cached",
    }
    assert '<c r="C1"><f>[1]Prices!B2*2</f><v>42</v></c>' in worksheet_xml(
        result["bytes"]
    )
    with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(result["bytes"])) as after:
        for name in before.namelist():
            if name.startswith("xl/externalLinks/"):
                assert after.read(name) == before.read(name)
    assert "external_links" not in fz.recalculate_xlsx_bytes(fixture_xlsx())


def test_external_links_policy_cached_or_refuse(tmp_path):
    source = external_link_fixture()
    cached = fz.recalculate_xlsx_bytes(source, external_links="cached")
    assert cached["external_links"]["policy"] == "cached"
    assert cached["bytes"] == fz.recalculate_xlsx_bytes(source)["bytes"]
    with pytest.raises(OSError, match="external link values"):
        fz.recalculate_xlsx_bytes(source, external_links="refuse")
    path = tmp_path / "linked.xlsx"
    path.write_bytes(source)
    with pytest.raises(OSError, match="external link values"):
        fz.recalculate_xlsx_file(str(path), external_links="refuse")
    assert path.read_bytes() == source
    # Links that nothing reads do not refuse.
    unused = fz.recalculate_xlsx_bytes(
        external_link_fixture(formula="1+2"), external_links="refuse"
    )
    assert "external_links" not in unused
    assert '<c r="C1"><f>1+2</f><v>3</v></c>' in worksheet_xml(unused["bytes"])
    for call in (
        lambda: fz.recalculate_xlsx_bytes(source, external_links="refresh"),
        lambda: fz.recalculate_xlsx_file(str(path), external_links="Cached"),
    ):
        with pytest.raises(ValueError, match="'cached' or 'refuse'"):
            call()
    assert path.read_bytes() == source
