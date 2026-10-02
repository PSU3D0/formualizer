import os
import subprocess
from io import BytesIO
from zipfile import ZipFile

import pytest

import formualizer as fz


@pytest.mark.parametrize("backend", ["python", "cli"])
@pytest.mark.parametrize("save_loop", [False, True])
def test_openpyxl_source_table_roundtrip(tmp_path, backend, save_loop):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.worksheet.table import Table

    path = tmp_path / "table.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Qty", "Price", "Amount"])
    sheet.append([2, 3, "=Table1[[#This Row],[Qty]]*Table1[[#This Row],[Price]]"])
    sheet.append([4, 5, "=[@Qty]*[@Price]"])
    sheet.add_table(Table(displayName="Table1", ref="A1:C3"))
    sheet["E1"] = "=SUM(Table1[Amount])"
    other = workbook.create_sheet("Other")
    other["A1"] = "=SUM(Table1[Qty])"
    workbook.save(path)
    if save_loop:
        workbook = openpyxl.load_workbook(path)
        workbook.active["A2"] = 7
        workbook.save(path)
    source = path.read_bytes()
    with ZipFile(BytesIO(source)) as archive:
        table = archive.read("xl/tables/table1.xml")
    if backend == "python":
        result = fz.recalculate_xlsx_bytes(source)
        path.write_bytes(result["bytes"])
    else:
        cli = os.environ.get("FORMUALIZER_CLI")
        if not cli:
            pytest.skip("set FORMUALIZER_CLI to test the candidate executable")
        result = subprocess.run([cli, "recalc", str(path)], capture_output=True)
        assert result.returncode == 0, result.stderr.decode()
    cached = openpyxl.load_workbook(path, data_only=True)
    assert cached.active["C2"].value == (21 if save_loop else 6)
    assert cached.active["C3"].value == 20
    assert cached.active["E1"].value == (41 if save_loop else 26)
    assert cached["Other"]["A1"].value == (11 if save_loop else 6)
    with ZipFile(path) as archive:
        assert archive.read("xl/tables/table1.xml") == table
    output = path.read_bytes()
    assert fz.recalculate_xlsx_bytes(output)["bytes"] == output


@pytest.mark.parametrize("backend", ["python", "cli"])
def test_table_formula_coverage_refuses_stale_extension(tmp_path, backend):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.worksheet.table import Table, TableColumn, TableFormula

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Qty", "Amount"])
    sheet.append([2, "=[@Qty]*10"])
    sheet.append([3, None])
    sheet.add_table(
        Table(
            displayName="Table1",
            ref="A1:B3",
            tableColumns=[
                TableColumn(id=1, name="Qty"),
                TableColumn(
                    id=2,
                    name="Amount",
                    calculatedColumnFormula=TableFormula(attr_text="[@Qty]*10"),
                ),
            ],
        )
    )
    path = tmp_path / "missing-formula.xlsx"
    workbook.save(path)
    original = path.read_bytes()
    if backend == "python":
        with pytest.raises(Exception, match="write the formula into each row"):
            fz.recalculate_xlsx_bytes(original)
    else:
        cli = os.environ.get("FORMUALIZER_CLI")
        if not cli:
            pytest.skip("set FORMUALIZER_CLI to test the candidate executable")
        result = subprocess.run([cli, "recalc", str(path)], capture_output=True)
        assert result.returncode == 2
        assert b"write the formula into each row" in result.stderr
        assert path.read_bytes() == original


@pytest.mark.parametrize("backend", ["python", "cli", "console"])
def test_native_source_today_is_the_host_date_not_the_epoch(backend, tmp_path):
    import datetime

    from openpyxl.utils.datetime import to_excel

    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    book.active["A1"] = "=TODAY()"
    stream = BytesIO()
    book.save(stream)
    before = to_excel(datetime.datetime.now().date())
    if backend == "python":
        out = fz.recalculate_xlsx_bytes(stream.getvalue())["bytes"]
    else:
        path = tmp_path / "today.xlsx"
        path.write_bytes(stream.getvalue())
        if backend == "console":
            from formualizer.cli import main

            assert main(["recalc", str(path)]) == 0
        else:
            cli = os.environ.get("FORMUALIZER_CLI")
            if not cli:
                pytest.skip("set FORMUALIZER_CLI to test the candidate executable")
            result = subprocess.run([cli, "recalc", str(path)], capture_output=True)
            assert result.returncode == 0, result.stderr.decode()
        out = path.read_bytes()
    after = to_excel(datetime.datetime.now().date())
    actual = openpyxl.load_workbook(BytesIO(out), data_only=True).active["A1"].value
    assert min(before, after) - 1 <= actual <= max(before, after) + 1
