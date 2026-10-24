import functools
import os

import pytest


def pytest_configure(config):
    """Exercise existing loader tests with an opt-in backend, never production defaults."""
    backend = os.environ.get("FORMUALIZER_TEST_BACKEND")
    if backend is None:
        return
    if backend not in {"calamine", "native"}:
        raise pytest.UsageError("FORMUALIZER_TEST_BACKEND must be calamine or native")
    import formualizer as fz

    patch = pytest.MonkeyPatch()
    config.add_cleanup(patch.undo)
    stats = {"backend": backend, "calls": 0}
    config._formualizer_backend_stats = stats

    def force(method, backend_position, class_method=False):
        @functools.wraps(method)
        def load(*args, **kwargs):
            args = list(args[1:] if class_method else args)
            positional = backend_position is not None and len(args) > backend_position
            selected = args[backend_position] if positional else kwargs.get("backend")
            # Explicit Umya coverage and explicit native tests remain intact.
            if selected is None or (
                isinstance(selected, str) and selected == "calamine"
            ):
                stats["calls"] += 1
                if positional:
                    args[backend_position] = backend
                else:
                    kwargs["backend"] = backend
            return method(*args, **kwargs)

        return classmethod(load) if class_method else load

    for name, position in [("from_path", 1), ("from_bytes", 1), ("load_path", 2)]:
        patch.setattr(
            fz.Workbook, name, force(getattr(fz.Workbook, name), position, True)
        )
    for name, position in [("load_workbook", None), ("load_workbook_bytes", 2)]:
        patch.setattr(fz, name, force(getattr(fz, name), position))


def pytest_terminal_summary(terminalreporter, exitstatus, config):
    stats = getattr(config, "_formualizer_backend_stats", None)
    if stats is not None:
        terminalreporter.write_line(
            f"Forced {stats['backend']} on {stats['calls']} default/Calamine load calls; "
            "explicit Umya coverage preserved."
        )


try:
    import openpyxl  # type: ignore
except Exception:  # pragma: no cover - allow skipping if not present in dev env
    openpyxl = None


pytestmark = pytest.mark.skipif(openpyxl is None, reason="openpyxl not installed")


@pytest.fixture
def xlsx_builder(tmp_path):
    """
    Returns a helper that creates a temporary XLSX at a unique path using openpyxl.
    The helper accepts a function (wb -> None) to populate the workbook before saving.
    """

    def _build(populate_fn):
        path = tmp_path / "wb.xlsx"
        wb = openpyxl.Workbook()
        # Ensure a predictable first sheet name
        wb.active.title = "Sheet1"
        populate_fn(wb)
        wb.save(path)
        return path

    return _build
