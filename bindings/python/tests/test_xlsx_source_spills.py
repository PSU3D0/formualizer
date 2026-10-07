"""Source-owned dynamic arrays: original OOXML, never an openpyxl save."""

from io import BytesIO
from stat import S_IMODE
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import openpyxl
import pytest

import formualizer as fz

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
DYNAMIC = "http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray"
NS = {"s": MAIN, "d": DYNAMIC}
SHEET = "xl/worksheets/sheet1.xml"
METADATA = (
    f'<metadata xmlns="{MAIN}" xmlns:xda="{DYNAMIC}">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" cellMeta="1"/></metadataTypes>'
    '<futureMetadata name="XLDAPR" count="1"><bk><extLst>'
    '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">'
    '<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/>'
    '</ext></extLst></bk></futureMetadata><cellMetadata count="1">'
    '<bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
)


def pack(parts):
    out = BytesIO()
    with ZipFile(out, "w") as archive:
        for name, body in parts.items():
            info = ZipInfo(name, (2020, 1, 2, 3, 4, 6))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, body)
    return out.getvalue()


def unpack(payload):
    with ZipFile(BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def fixture(
    *,
    old=0,
    size=3,
    formula="_xlfn.SEQUENCE(Size)",
    blocker=False,
    epoch=False,
    merge=False,
    cse=False,
):
    attrs = ' cm="1"' if old else ""
    fattrs = f' t="array" ref="C2:C{old + 1}"' if old else ""
    if cse:
        fattrs = ' t="array" ref="C2:C4"'
    rows = '<row r="1"><c r="B1"><v>' + str(size) + "</v></c></row>"
    rows += f'<row r="2"><c r="C2"{attrs}><f{fattrs}>{formula}</f><v>99</v></c></row>'
    for row in range(3, max(old + 2, 7 if blocker else 3)):
        value = 123 if blocker and row == 6 else row - 1
        if row <= old + 1 or (blocker and row == 6):
            rows += f'<row r="{row}"><c r="C{row}" s="1"><v>{value}</v></c></row>'
    rows += '<row r="9"><c r="C9"><f>SUM(C2#)</f><v>99</v></c></row>'
    rows += (
        '<row r="10"><c r="C10"><f>SUM(_xlfn.ANCHORARRAY(C2))</f><v>99</v></c></row>'
    )
    types = f'<Types xmlns="{CT}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="bin" ContentType="application/octet-stream"/>'
    for part, kind in [
        ("workbook", "sheet.main"),
        ("worksheets/sheet1", "worksheet"),
        ("styles", "styles"),
    ]:
        types += f'<Override PartName="/xl/{part}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.{kind}+xml"/>'
    rels = f'<Relationships xmlns="{RELS}"><Relationship Id="rId1" Type="{OFFICE}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{OFFICE}/styles" Target="styles.xml"/>'
    if old:
        types += '<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>'
        rels += f'<Relationship Id="rId3" Type="{OFFICE}/sheetMetadata" Target="metadata.xml"/>'
    parts = {
        "[Content_Types].xml": types + "</Types>",
        "_rels/.rels": f'<Relationships xmlns="{RELS}"><Relationship Id="rId1" Type="{OFFICE}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{MAIN}" xmlns:r="{OFFICE}"><workbookPr date1904="{int(epoch)}"/><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets><definedNames><definedName name="Size">Sheet1!$B$1</definedName></definedNames></workbook>',
        "xl/_rels/workbook.xml.rels": rels + "</Relationships>",
        "xl/styles.xml": f'<styleSheet xmlns="{MAIN}"><fonts count="2"><font/><font><b/></font></fonts><fills count="1"><fill><patternFill/></fill></fills><borders count="1"><border/></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="2"><xf numFmtId="0" fontId="0"/><xf numFmtId="0" fontId="1" applyFont="1"/></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>',
        SHEET: f'<worksheet xmlns="{MAIN}"><dimension ref="B1:C10"/><sheetData>{rows}</sheetData>'
        + (
            '<mergeCells count="1"><mergeCell ref="C3:D3"/></mergeCells>'
            if merge
            else ""
        )
        + "</worksheet>",
        "custom/opaque.bin": b"opaque\x00payload\xff",
    }
    if old:
        parts["xl/metadata.xml"] = METADATA
    return pack(parts)


def patch(payload, old, new):
    parts = unpack(payload)
    assert old.encode() in parts[SHEET]
    parts[SHEET] = parts[SHEET].replace(old.encode(), new.encode(), 1)
    return pack(parts)


def inspect(payload, formula, extent):
    parts = unpack(payload)
    cells = {c.get("r"): c for c in ET.fromstring(parts[SHEET]).findall(".//s:c", NS)}
    anchor = cells["C2"]
    f = anchor.find("s:f", NS)
    assert f.text == formula
    assert cells["C9"].findtext("s:f", namespaces=NS) == "SUM(C2#)"
    assert cells["C10"].findtext("s:f", namespaces=NS) == "SUM(_xlfn.ANCHORARRAY(C2))"
    assert f.get("t") == "array" and f.get("ref") == extent
    metadata = ET.fromstring(parts["xl/metadata.xml"])
    block = metadata.findall("s:cellMetadata/s:bk", NS)[int(anchor.get("cm")) - 1]
    rc = block.find("s:rc", NS)
    assert (
        metadata.findall("s:metadataTypes/s:metadataType", NS)[
            int(rc.get("t")) - 1
        ].get("name")
        == "XLDAPR"
    )
    future = metadata.findall('s:futureMetadata[@name="XLDAPR"]/s:bk', NS)[
        int(rc.get("v"))
    ]
    ext = future.find("s:extLst/s:ext", NS)
    assert ext.get("uri") == "{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"
    assert ext.find("d:dynamicArrayProperties", NS).get("fDynamic") == "1"
    rel = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    assert (
        len(
            [
                r
                for r in rel
                if r.get("Type") == OFFICE + "/sheetMetadata"
                and r.get("Target") == "metadata.xml"
            ]
        )
        == 1
    )
    assert any(
        c.get("PartName") == "/xl/metadata.xml"
        for c in ET.fromstring(parts["[Content_Types].xml"])
    )
    assert parts["custom/opaque.bin"] == b"opaque\x00payload\xff"
    return cells


def compressed_member(payload, name):
    with ZipFile(BytesIO(payload)) as archive:
        info = archive.getinfo(name)
        offset = info.header_offset
        # ZIP32 local header: filename/extra lengths at offsets 26 and 28.
        start = (
            offset + 30 + int.from_bytes(payload[offset + 26 : offset + 28], "little")
        )
        start += int.from_bytes(payload[offset + 28 : offset + 30], "little")
        return payload[start : start + info.compress_size]


def exercise(tmp_path, payload):
    result = fz.recalculate_xlsx_bytes(payload)
    before, after = unpack(payload), unpack(result["bytes"])
    for name in ("xl/workbook.xml", "xl/styles.xml", "custom/opaque.bin"):
        assert before[name] == after[name]
        assert compressed_member(payload, name) == compressed_member(
            result["bytes"], name
        )
    source, destination = tmp_path / "source.xlsx", tmp_path / "destination.xlsx"
    source.write_bytes(payload)
    destination.write_bytes(b"previous destination")
    destination.chmod(0o640)
    file_result = fz.recalculate_xlsx_file(str(source), output=str(destination))
    assert destination.read_bytes() == result["bytes"]
    assert S_IMODE(destination.stat().st_mode) == 0o640
    for key in ("formula_cells", "cache_cells_changed", "worksheet_parts_changed"):
        assert file_result[key] == result[key]
    assert source.read_bytes() == payload
    fresh = tmp_path / "new.xlsx"
    fz.recalculate_xlsx_file(str(source), output=str(fresh))
    assert fresh.read_bytes() == result["bytes"]
    fz.recalculate_xlsx_file(str(source))
    assert source.read_bytes() == result["bytes"]
    repeated = fz.recalculate_xlsx_bytes(result["bytes"])
    assert repeated["bytes"] == result["bytes"]
    assert repeated["cache_cells_changed"] == 0
    assert repeated["worksheet_parts_changed"] == 0
    return result


@pytest.mark.parametrize(
    "old,size,changed", [(0, 3, 5), (3, 5, 5), (5, 2, 6), (5, 1, 7)]
)
def test_new_grow_shrink_collapse(tmp_path, old, size, changed):
    result = exercise(tmp_path, fixture(old=old, size=size))
    assert result["formula_cells"] == result["summary"]["evaluated"] == 3
    assert result["cache_cells_changed"] == changed
    assert result["worksheet_parts_changed"] == 1
    cells = inspect(
        result["bytes"],
        "_xlfn.SEQUENCE(Size)",
        "C2" if size == 1 else f"C2:C{size + 1}",
    )
    for row in range(size + 2, old + 2):
        assert cells[f"C{row}"].get("s") == "1"
        assert cells[f"C{row}"].find("s:v", NS) is None
    book = openpyxl.load_workbook(BytesIO(result["bytes"]), data_only=True)
    assert [book.active.cell(row, 3).value for row in range(2, size + 2)] == list(
        range(1, size + 1)
    )
    assert book.active["C9"].value == book.active["C10"].value == size * (size + 1) / 2
    assert all(
        book.active.cell(row, 3).value is None for row in range(size + 2, old + 2)
    )
    assert all(book.active.cell(row, 3).font.bold for row in range(size + 2, old + 2))


def test_blocked_then_source_xml_reexpansion(tmp_path):
    result = exercise(tmp_path, fixture(old=3, size=5, blocker=True))
    inspect(result["bytes"], "_xlfn.SEQUENCE(Size)", "C2")
    book = openpyxl.load_workbook(BytesIO(result["bytes"]), data_only=True)
    assert book.active["C2"].value == "#SPILL!"
    assert book.active["C6"].value == 123
    assert book.active["C3"].value is None
    assert book.active["C9"].value == book.active["C10"].value == "#REF!"
    payload = patch(
        result["bytes"], '<c r="C6" s="1"><v>123</v></c>', '<c r="C6" s="1"></c>'
    )
    expanded = exercise(tmp_path, payload)
    inspect(expanded["bytes"], "_xlfn.SEQUENCE(Size)", "C2:C6")
    book = openpyxl.load_workbook(BytesIO(expanded["bytes"]), data_only=True)
    assert book.active["C6"].value == 5
    assert book.active["C9"].value == book.active["C10"].value == 15


@pytest.mark.parametrize(
    "formula,values",
    [
        ("{1;2}", [1, 2]),
        ("{&quot;a&quot;;&quot;b&quot;}", ["a", "b"]),
        ("{TRUE;FALSE}", [True, False]),
        ("{#N/A;#DIV/0!}", ["#N/A", "#DIV/0!"]),
    ],
)
def test_typed_members(tmp_path, formula, values):
    result = exercise(tmp_path, fixture(formula=formula))
    cells = inspect(result["bytes"], ET.fromstring(f"<f>{formula}</f>").text, "C2:C3")
    book = openpyxl.load_workbook(BytesIO(result["bytes"]), data_only=True)
    assert [book.active["C2"].value, book.active["C3"].value] == values
    expected_type = (
        "b"
        if isinstance(values[0], bool)
        else "e"
        if str(values[0]).startswith("#")
        else "str"
        if isinstance(values[0], str)
        else None
    )
    assert cells["C3"].get("t") == expected_type


def test_source_1904_epoch(tmp_path):
    formula = "DATE(2020,1,1)+SEQUENCE(2,1,0)"
    result = exercise(tmp_path, fixture(formula=formula, epoch=True))
    inspect(result["bytes"], formula, "C2:C3")
    book = openpyxl.load_workbook(BytesIO(result["bytes"]), data_only=True)
    assert book.active["C2"].value == 42369
    assert book.active["C3"].value == 42370


@pytest.mark.parametrize("kwargs,message", [({"merge": True}, "merged")])
def test_refusal_preserves_destination_bytes_and_permissions(tmp_path, kwargs, message):
    payload = fixture(**kwargs)
    with pytest.raises(OSError, match=message):
        fz.recalculate_xlsx_bytes(payload)
    source, destination = tmp_path / "source.xlsx", tmp_path / "destination.xlsx"
    source.write_bytes(payload)
    destination.write_bytes(b"do not publish")
    destination.chmod(0o640)
    with pytest.raises(OSError, match=message):
        fz.recalculate_xlsx_file(str(source), output=str(destination))
    assert destination.read_bytes() == b"do not publish"
    assert S_IMODE(destination.stat().st_mode) == 0o640
    assert source.read_bytes() == payload
