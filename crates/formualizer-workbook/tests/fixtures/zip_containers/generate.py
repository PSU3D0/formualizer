# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Regenerate the ZIP-container fixtures for tests/xlsx_zip_containers.rs.

Every fixture holds the same openpyxl-made workbook (formulas without cached
values, so recalculation writes) in the container shape a real producer
writes:

- openpyxl.xlsx: openpyxl itself (no extra fields, no data descriptors).
- excel_growth_hint.xlsx: re-containered in Excel's layout: version made by
  4.5, needed 2.0, flags 0x0006, deflate, DOS time 1980-01-01 00:00, and the
  Microsoft 0xA220 growth-hint extra field (signature 0xA028, the padding
  size, zero padding) in the local headers of the package parts Excel pads
  (512 bytes for [Content_Types].xml and _rels/.rels, 256 for the other
  relationship and docProps parts); central records carry no extras.
- descriptor_unsigned.xlsx: data descriptors without the optional signature,
  zero CRC/sizes in the local headers.
- ntfs_times.xlsx: NTFS (0x000A) times in local and central records.
- info_zip.xlsx: Info-ZIP `zip -r` (extended timestamp 0x5455 and Unix
  UID/GID 0x7875 extra fields, directory entries).
- libreoffice.xlsx: LibreOffice headless conversion (signed data
  descriptors).

Requires `zip` and `soffice` on PATH. Run with `uv run generate.py`.
"""
import os
import pathlib
import shutil
import struct
import subprocess
import tempfile
import zipfile
import zlib

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent


def workbook(path):
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    for row in range(1, 6):
        data.cell(row, 1, row * 10)
        data.cell(row, 2, f"=A{row}*2")
    data["C1"] = "label"
    summary = wb.create_sheet("Summary")
    summary["A1"] = "=SUM(Data!B1:B5)"
    summary["A2"] = '=Data!C1&" total"'
    wb.save(path)


def members(path):
    with zipfile.ZipFile(path) as z:
        return [(i.filename, z.read(i.filename)) for i in z.infolist()]


def deflate(data):
    c = zlib.compressobj(1, zlib.DEFLATED, -15)
    return c.compress(data) + c.flush()


def growth_hint(pad):
    return struct.pack("<HHHH", 0xA220, 4 + pad, 0xA028, pad) + bytes(pad)


def ntfs(seed):
    times = struct.pack("<QQQ", 132000000000000000 + seed, 132000000000000000, 132000000000000000)
    return struct.pack("<HHIHH", 0x000A, 32, 0, 1, 24) + times


def container(parts, made_by, flags, local_extra, central_extra, descriptor):
    """Raw ZIP32 writer: one local record, payload and optional descriptor
    per part, then the central directory and the end record."""
    out = bytearray()
    central = bytearray()
    for index, (name, data) in enumerate(parts):
        raw = name.encode()
        body = deflate(data)
        crc = zlib.crc32(data)
        sizes = struct.pack("<III", crc, len(body), len(data))
        lx, cx = local_extra(name, index), central_extra(name, index)
        offset = len(out)
        fixed = struct.pack("<HHHHH", 20, flags, 8, 0, 0x0021)
        local_sizes = bytes(12) if descriptor else sizes
        out += b"PK\x03\x04" + fixed + local_sizes + struct.pack("<HH", len(raw), len(lx)) + raw + lx + body
        if descriptor == "unsigned":
            out += sizes
        central += (
            b"PK\x01\x02"
            + struct.pack("<H", made_by)
            + fixed
            + sizes
            + struct.pack("<HHHHHII", len(raw), len(cx), 0, 0, 0, 0, offset)
            + raw
            + cx
        )
    start = len(out)
    out += central
    out += b"PK\x05\x06" + struct.pack("<HHHHIIH", 0, 0, len(parts), len(parts), len(central), start, 0)
    return bytes(out)


EXCEL_PADS = {"[Content_Types].xml": 512, "_rels/.rels": 512}


def excel_pad(name, _):
    if name in EXCEL_PADS:
        return growth_hint(EXCEL_PADS[name])
    if name.endswith(".rels") or name.startswith("docProps/"):
        return growth_hint(256)
    return b""


def main():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = pathlib.Path(tmp)
        base = tmp / "openpyxl.xlsx"
        workbook(base)
        shutil.copy(base, HERE / "openpyxl.xlsx")
        parts = members(base)
        none = lambda *_: b""
        (HERE / "excel_growth_hint.xlsx").write_bytes(container(parts, 45, 0x0006, excel_pad, none, None))
        (HERE / "descriptor_unsigned.xlsx").write_bytes(container(parts, 20, 0x0008, none, none, "unsigned"))
        times = lambda _, i: ntfs(i)
        (HERE / "ntfs_times.xlsx").write_bytes(container(parts, 10 << 8 | 63, 0, times, times, None))

        tree = tmp / "tree"
        with zipfile.ZipFile(base) as z:
            z.extractall(tree)
        for path in sorted(tree.rglob("*"), reverse=True):
            os.utime(path, (1577934245, 1577934245))
        os.utime(tree, (1577934245, 1577934245))
        target = HERE / "info_zip.xlsx"
        target.unlink(missing_ok=True)
        names = [n for n, _ in parts]
        top = sorted({n.split("/")[0] for n in names}, key=lambda n: (n != "[Content_Types].xml", n))
        subprocess.run(["zip", "-q", "-r", str(target), *top], cwd=tree, check=True, env=dict(os.environ, TZ="UTC"))

        profile = tmp / "profile"
        out = tmp / "lo"
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(out),
                str(base),
            ],
            check=True,
            capture_output=True,
        )
        shutil.copy(out / "openpyxl.xlsx", HERE / "libreoffice.xlsx")


if __name__ == "__main__":
    main()
