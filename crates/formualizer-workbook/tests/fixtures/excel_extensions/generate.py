# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Regenerate the Excel-extension fixtures for tests/xlsx_excel_extensions.rs.

Each fixture is an openpyxl workbook into which the extension markup that
Excel 2010+ writes routinely is copied, in Excel's layout:

- date1904_x15.xlsx / date1900_x15.xlsx: a 1904 (resp. 1900) workbook whose
  workbook.xml carries Excel 2013+'s root namespaces, `mc:AlternateContent`
  `x15ac:absPath`, `xr:revisionPtr` and the workbook `extLst` with
  `<x15:workbookPr chartTrackingRefBase="1"/>` and `xcalcf:calcFeatures`.
- controls.xlsx: a worksheet with an Excel 2010 form control (`controls`
  inside `mc:AlternateContent`, anchored with `xdr:col`/`xdr:row`) and an
  x14 list data validation (`xm:f`, `xm:sqref`) in the worksheet `extLst`,
  plus the control's ctrlProp part. controls_plain.xlsx is the same
  workbook without that markup.
- tables_xr.xlsx: a table with Excel 2016+'s `mc:Ignorable="xr xr3"`,
  `xr:uid` (table, autoFilter) and `xr3:uid` (table columns).

openpyxl writes formulas without cached values, so recalculation writes.
Run with `uv run generate.py`.
"""
import datetime
import pathlib
import zipfile

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = pathlib.Path(__file__).resolve().parent
STAMP = datetime.datetime(2024, 1, 2, 3, 4, 5)

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
X14AC = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"
X15 = "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"
X15AC = "http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
XR = "http://schemas.microsoft.com/office/spreadsheetml/2014/revision"
XR2 = "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2"
XR3 = "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3"
XR6 = "http://schemas.microsoft.com/office/spreadsheetml/2016/revision6"
XR10 = "http://schemas.microsoft.com/office/spreadsheetml/2016/revision10"
XCALCF = "http://schemas.microsoft.com/office/spreadsheetml/2018/calcfeatures"
MAIN_DECL = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


def stamp(wb):
    wb.properties.created = STAMP
    wb.properties.modified = STAMP


def save_with(src, dst, edits, extra=()):
    """Copy `src` to `dst`, applying `edits[name](text) -> text` to members
    and appending `extra` (name, text) members, in openpyxl's container."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename in edits:
                data = edits[info.filename](data.decode()).encode()
            out = zipfile.ZipInfo(info.filename, date_time=STAMP.timetuple()[:6])
            out.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(out, data)
        for name, text in extra:
            out = zipfile.ZipInfo(name, date_time=STAMP.timetuple()[:6])
            out.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(out, text.encode())


def replace_once(text, old, new):
    assert text.count(old) == 1, (old, text[:300])
    return text.replace(old, new)


def excel_workbook_xml(text):
    """Excel 2013+/365 workbook.xml layout around openpyxl's content."""
    text = replace_once(
        text,
        MAIN_DECL,
        f'{MAIN_DECL} xmlns:mc="{MC}" mc:Ignorable="x15 xr xr6 xr10 xr2" xmlns:x15="{X15}"'
        f' xmlns:xr="{XR}" xmlns:xr6="{XR6}" xmlns:xr10="{XR10}" xmlns:xr2="{XR2}"',
    )
    absolute = (
        f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice Requires="x15">'
        f'<x15ac:absPath url="C:\\Users\\fixture\\" xmlns:x15ac="{X15AC}"/>'
        "</mc:Choice></mc:AlternateContent>"
        '<xr:revisionPtr revIDLastSave="0" documentId="8_{00000000-0000-0000-0000-000000000001}"'
        ' xr6:coauthVersionLast="47" xr6:coauthVersionMax="47"'
        ' xr10:uidLastSave="{00000000-0000-0000-0000-000000000000}"/>'
    )
    text = replace_once(text, "<bookViews>", absolute + "<bookViews>")
    ext = (
        "<extLst>"
        '<ext uri="{140A7094-0E35-4892-8432-C4D2E57EDEB5}" xmlns:x15="' + X15 + '">'
        '<x15:workbookPr chartTrackingRefBase="1"/></ext>'
        '<ext uri="{B58B0392-4F1F-4190-BB64-5DF3571DCE5F}" xmlns:xcalcf="' + XCALCF + '">'
        '<xcalcf:calcFeatures><xcalcf:feature name="microsoft.com:RD"/>'
        '<xcalcf:feature name="microsoft.com:Single"/><xcalcf:feature name="microsoft.com:FV"/>'
        '<xcalcf:feature name="microsoft.com:CNMTM"/><xcalcf:feature name="microsoft.com:LET_WF"/>'
        '<xcalcf:feature name="microsoft.com:LAMBDA_WF"/><xcalcf:feature name="microsoft.com:ARRAYTEXT_WF"/>'
        "</xcalcf:calcFeatures></ext>"
        "</extLst>"
    )
    return replace_once(text, "</workbook>", ext + "</workbook>")


def dates(path, epoch_1904):
    wb = openpyxl.Workbook()
    stamp(wb)
    if epoch_1904:
        wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Dates"
    ws["A1"] = datetime.datetime(2024, 1, 15)
    ws["A1"].number_format = "yyyy-mm-dd"
    ws["A2"] = datetime.datetime(1904, 1, 2)
    ws["A2"].number_format = "yyyy-mm-dd"
    ws["B1"] = "=YEAR(A1)"
    ws["B2"] = "=YEAR(A2)"
    ws["C1"] = "=A1+1"
    ws["C1"].number_format = "yyyy-mm-dd"
    ws["D1"] = "=DATE(2024,1,15)"
    ws["D2"] = "=DATE(1904,1,2)"
    ws["E1"] = "=DAY(A1)"
    wb.save(path)


def controls_workbook(path):
    wb = openpyxl.Workbook()
    stamp(wb)
    ws = wb.active
    ws.title = "Data"
    for row in range(1, 6):
        ws.cell(row, 1, row * 10)
        ws.cell(row, 2, f"=A{row}*2")
    ws["C1"] = "beta"
    ws["D1"] = "=SUM(B1:B5)"
    ws["D2"] = '=MATCH(C1,Lists!A1:A3,0)'
    ws["D3"] = "=ROWS(A1:A5)+COUNT(A:A)"
    lists = wb.create_sheet("Lists")
    for row, name in enumerate(["alpha", "beta", "gamma"], 1):
        lists.cell(row, 1, name)
    wb.save(path)


def excel_controls_sheet(text):
    text = replace_once(
        text,
        MAIN_DECL,
        f'{MAIN_DECL} xmlns:r="{R}" xmlns:xdr="{XDR}" xmlns:x14="{X14}" xmlns:mc="{MC}"'
        f' mc:Ignorable="x14ac xr xr2 xr3" xmlns:x14ac="{X14AC}" xmlns:xr="{XR}"'
        f' xmlns:xr2="{XR2}" xmlns:xr3="{XR3}" xr:uid="{{00000000-0001-0000-0000-000000000000}}"',
    )
    anchor = (
        '<anchor moveWithCells="1" sizeWithCells="1">'
        "<from><xdr:col>5</xdr:col><xdr:colOff>228600</xdr:colOff><xdr:row>0</xdr:row>"
        "<xdr:rowOff>139700</xdr:rowOff></from>"
        "<to><xdr:col>7</xdr:col><xdr:colOff>6350</xdr:colOff><xdr:row>3</xdr:row>"
        "<xdr:rowOff>63500</xdr:rowOff></to></anchor>"
    )
    controls = (
        f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice Requires="x14"><controls>'
        f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice Requires="x14">'
        '<control shapeId="1025" r:id="rId1" name="Button 1">'
        '<controlPr defaultSize="0" print="0" autoFill="0" autoPict="0" macro="[0]!Button1_Click">'
        + anchor
        + "</controlPr></control></mc:Choice></mc:AlternateContent>"
        "</controls></mc:Choice></mc:AlternateContent>"
    )
    validation = (
        f'<extLst><ext uri="{{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}}" xmlns:x14="{X14}">'
        f'<x14:dataValidations count="1" xmlns:xm="{XM}">'
        '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" showErrorMessage="1"'
        ' xr:uid="{00000000-0002-0000-0000-000000000000}">'
        "<x14:formula1><xm:f>Lists!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>C1</xm:sqref>"
        "</x14:dataValidation></x14:dataValidations></ext></extLst>"
    )
    return replace_once(text, "</worksheet>", controls + validation + "</worksheet>")


def add_ctrl_prop(text):
    return replace_once(
        text,
        "</Types>",
        '<Override PartName="/xl/ctrlProps/ctrlProp1.xml"'
        ' ContentType="application/vnd.ms-excel.controlproperties+xml"/></Types>',
    )


def tables(path):
    wb = openpyxl.Workbook()
    stamp(wb)
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Qty", "Price"])
    for qty, price in [(2, 1.5), (3, 2.5), (4, 3.5)]:
        ws.append([qty, price])
    ws["D1"] = "=SUM(Sales[Qty])"
    ws["D2"] = "=SUMPRODUCT(Sales[Qty],Sales[Price])"
    table = Table(displayName="Sales", ref="A1:B4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    wb.save(path)


def excel_table_xml(text):
    """The same table in Excel 2016+'s layout (copied from an Excel-saved part)."""
    assert 'name="Sales"' in text and 'ref="A1:B4"' in text, text
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:mc="{MC}"'
        f' mc:Ignorable="xr xr3" xmlns:xr="{XR}" xmlns:xr3="{XR3}" id="1"'
        ' xr:uid="{00000000-000C-0000-FFFF-FFFF00000000}" name="Sales" displayName="Sales"'
        ' ref="A1:B4" totalsRowShown="0">'
        '<autoFilter ref="A1:B4" xr:uid="{00000000-0009-0000-0100-000001000000}"/>'
        '<tableColumns count="2">'
        '<tableColumn id="1" xr3:uid="{00000000-0010-0000-0000-000001000000}" name="Qty"/>'
        '<tableColumn id="2" xr3:uid="{00000000-0010-0000-0000-000002000000}" name="Price"/>'
        "</tableColumns>"
        '<tableStyleInfo name="TableStyleMedium2" showFirstColumn="0" showLastColumn="0"'
        ' showRowStripes="1" showColumnStripes="0"/></table>'
    )


def main():
    for name, epoch in [("date1904_x15.xlsx", True), ("date1900_x15.xlsx", False)]:
        base = HERE / ("_" + name)
        dates(base, epoch)
        save_with(base, HERE / name, {"xl/workbook.xml": excel_workbook_xml})
        base.unlink()

    plain = HERE / "controls_plain.xlsx"
    controls_workbook(plain)
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2006/relationships/ctrlProp"'
        ' Target="../ctrlProps/ctrlProp1.xml"/></Relationships>'
    )
    ctrl = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<formControlPr xmlns="{X14}" objectType="Button" lockText="1"/>'
    )
    save_with(
        plain,
        HERE / "controls.xlsx",
        {"xl/worksheets/sheet1.xml": excel_controls_sheet, "[Content_Types].xml": add_ctrl_prop},
        [("xl/worksheets/_rels/sheet1.xml.rels", rels), ("xl/ctrlProps/ctrlProp1.xml", ctrl)],
    )

    base = HERE / "_tables.xlsx"
    tables(base)
    with zipfile.ZipFile(base) as z:
        table_part = next(n for n in z.namelist() if n.startswith("xl/tables/"))
    save_with(base, HERE / "tables_xr.xlsx", {table_part: excel_table_xml})
    base.unlink()


if __name__ == "__main__":
    main()
